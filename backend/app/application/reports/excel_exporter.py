"""Exportación de estados financieros a Excel con formato corporativo VANTAGE."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

from app.application.report_schemas import BalanceSheetResponse, IncomeStatementResponse

VIOLET = "7C3AED"
VIOLET_DARK = "5B21B6"
VIOLET_PALE = "F5F3FF"
SOFT_BG = "F9FAFB"
BORDER = "E5E7EB"
TOTAL_BG = "E5E7EB"
DARK = "111827"

MONEY_FORMAT = "#,##0.00"
HEADERS = ("Tipo", "Código", "Cuenta", "Saldo")
THIN_BORDER = Border(bottom=Side(style="thin", color=BORDER))
TOTAL_TOP_BORDER = Border(
    top=Side(style="thin", color=VIOLET_DARK),
    bottom=Side(style="thin", color=BORDER),
)


def _decimal(value: str | Decimal | int) -> Decimal:
    return Decimal(str(value))


def _period_label(from_date: date | None, to_date: date | None) -> str:
    if from_date is not None and to_date is not None:
        return f"Periodo: {from_date.strftime('%d/%m/%Y')} - {to_date.strftime('%d/%m/%Y')}"
    if to_date is not None:
        return f"Corte: {to_date.strftime('%d/%m/%Y')}"
    if from_date is not None:
        return f"Desde: {from_date.strftime('%d/%m/%Y')}"
    return "Periodo completo"


def _write_report(
    *,
    title: str,
    sheet_title: str,
    subtitle: str,
    rows: list[tuple[str, str, str, Decimal]],
    totals: list[tuple[str, str, str, Decimal]],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.sheet_properties.tabColor = VIOLET
    sheet.sheet_view.showGridLines = False

    # Fila 1: título corporativo.
    sheet.merge_cells("A1:D1")
    title_cell = sheet["A1"]
    title_cell.value = title
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor=VIOLET)
    title_cell.alignment = Alignment(vertical="center", horizontal="left")
    sheet.row_dimensions[1].height = 30

    # Fila 2: empresa y periodo.
    sheet.merge_cells("A2:D2")
    subtitle_cell = sheet["A2"]
    subtitle_cell.value = subtitle
    subtitle_cell.font = Font(name="Calibri", size=10, color=DARK)
    subtitle_cell.fill = PatternFill("solid", fgColor=VIOLET_PALE)
    subtitle_cell.alignment = Alignment(vertical="center", horizontal="left")
    sheet.row_dimensions[2].height = 22

    # Fila 4: encabezados.
    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font = Font(name="Calibri", size=10, bold=True, color=DARK)
        cell.fill = PatternFill("solid", fgColor="F3F4F6")
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center")

    # Filas de datos.
    first_data_row = 5
    for index, row in enumerate(rows):
        excel_row = first_data_row + index
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=excel_row, column=column, value=value)
            cell.font = Font(name="Calibri", size=10, color=DARK)
            cell.border = THIN_BORDER
            if index % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=SOFT_BG)
        sheet.cell(row=excel_row, column=4).number_format = MONEY_FORMAT
        sheet.cell(row=excel_row, column=4).alignment = Alignment(horizontal="right")

    # Filas de totales.
    total_start = first_data_row + len(rows)
    for index, row in enumerate(totals):
        excel_row = total_start + index
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=excel_row, column=column, value=value)
            cell.font = Font(name="Calibri", size=10, bold=True, color=DARK)
            cell.fill = PatternFill("solid", fgColor=TOTAL_BG)
            cell.border = TOTAL_TOP_BORDER if index == 0 else THIN_BORDER
            if column == 4:
                cell.number_format = MONEY_FORMAT
                cell.alignment = Alignment(horizontal="right")

    # Anchura automática con topes razonables.
    widths = {1: 16, 2: 14, 3: 48, 4: 18}
    last_row = total_start + len(totals)
    for column in range(1, 5):
        max_length = widths[column]
        for row in sheet.iter_rows(
            min_row=4,
            max_row=last_row,
            min_col=column,
            max_col=column,
        ):
            value = row[0].value
            if value is not None:
                max_length = max(max_length, len(str(value)))
        max_width = 60 if column == 3 else 22
        sheet.column_dimensions[chr(64 + column)].width = min(max_length + 2, max_width)

    sheet.freeze_panes = "A5"
    sheet.print_title_rows = "1:4"
    sheet.auto_filter.ref = f"A4:D{last_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_income_statement_xlsx(
    statement: IncomeStatementResponse,
    *,
    company_name: str,
    from_date: date | None,
    to_date: date | None,
) -> bytes:
    """Genera el Estado de Resultados en XLSX con estilo VANTAGE."""
    rows = [
        ("Ingreso", account.code, account.name, _decimal(account.balance))
        for account in statement.income_accounts
    ]
    rows.extend(
        ("Gasto/Costo", account.code, account.name, _decimal(account.balance))
        for account in statement.expense_accounts
    )
    totals = [
        ("TOTAL INGRESOS", "", "", _decimal(statement.total_income)),
        ("TOTAL GASTOS Y COSTOS", "", "", _decimal(statement.total_expenses)),
        ("UTILIDAD NETA", "", "", _decimal(statement.net_income)),
    ]
    return _write_report(
        title="VANTAGE - Estado de Resultados",
        sheet_title="Estado de Resultados",
        subtitle=f"Empresa: {company_name}  ·  {_period_label(from_date, to_date)}",
        rows=rows,
        totals=totals,
    )


def export_balance_sheet_xlsx(
    statement: BalanceSheetResponse,
    *,
    company_name: str,
    as_of_date: date | None,
) -> bytes:
    """Genera el Balance General en XLSX con estilo VANTAGE."""
    rows = [
        ("Activo", account.code, account.name, _decimal(account.balance))
        for account in statement.asset_accounts
    ]
    rows.extend(
        ("Pasivo", account.code, account.name, _decimal(account.balance))
        for account in statement.liability_accounts
    )
    equity_account_total = sum(
        (_decimal(account.balance) for account in statement.equity_accounts),
        Decimal("0.00"),
    )
    rows.extend(
        ("Patrimonio", account.code, account.name, _decimal(account.balance))
        for account in statement.equity_accounts
    )
    period_result = _decimal(statement.total_equity) - equity_account_total
    if period_result != Decimal("0.00"):
        rows.append(("Patrimonio", "", "Resultado del ejercicio", period_result))

    totals = [
        ("TOTAL ACTIVO", "", "", _decimal(statement.total_assets)),
        ("TOTAL PASIVO", "", "", _decimal(statement.total_liabilities)),
        ("TOTAL PATRIMONIO", "", "", _decimal(statement.total_equity)),
        (
            "PASIVO + PATRIMONIO",
            "",
            "",
            _decimal(statement.total_liabilities) + _decimal(statement.total_equity),
        ),
    ]
    return _write_report(
        title="VANTAGE - Balance General",
        sheet_title="Balance General",
        subtitle=f"Empresa: {company_name}  ·  {_period_label(None, as_of_date)}",
        rows=rows,
        totals=totals,
    )


__all__ = ["export_balance_sheet_xlsx", "export_income_statement_xlsx"]
