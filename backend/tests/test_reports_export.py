"""Pruebas de exportación XLSX de los estados financieros."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.domain.models.accounting import (
    AccountingEntry,
    AccountingEntryLine,
    AccountingEntryStatus,
    AccountingSourceType,
)

D = Decimal
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _post_entry(
    db: Session,
    entry_date: date,
    description: str,
    lines: list[tuple[str, Decimal, Decimal]],
) -> None:
    entry = AccountingEntry(
        entry_date=entry_date,
        description=description,
        source_type=AccountingSourceType.EXPENSE,
        source_id=uuid4(),
        status=AccountingEntryStatus.POSTED,
    )
    db.add(entry)
    db.flush()
    for code, debit, credit in lines:
        db.add(
            AccountingEntryLine(
                entry_id=entry.id,
                account_code=code,
                debit=debit,
                credit=credit,
            )
        )
    db.flush()


def test_export_income_statement_xlsx(
    client: TestClient,
    db_session: Session,
) -> None:
    _post_entry(
        db_session,
        date(2026, 3, 2),
        "Ingreso exportación",
        [("110505", D("1000.00"), D("0.00")), ("4135", D("0.00"), D("1000.00"))],
    )
    _post_entry(
        db_session,
        date(2026, 3, 3),
        "Gasto exportación",
        [("510506", D("250.00"), D("0.00")), ("110505", D("0.00"), D("250.00"))],
    )

    response = client.get(
        "/api/v1/reports/export/xlsx",
        params={
            "report": "income-statement",
            "from_date": "2026-03-01",
            "to_date": "2026-03-31",
        },
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert 'filename="estado-resultados.xlsx"' in response.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    sheet = workbook.active
    assert sheet["A1"].value == "VANTAGE - Estado de Resultados"
    assert "Empresa:" in str(sheet["A2"].value)
    assert [sheet.cell(row=4, column=column).value for column in range(1, 5)] == [
        "Tipo",
        "Código",
        "Cuenta",
        "Saldo",
    ]

    values: dict[str, Decimal] = {}
    number_formats: dict[str, str] = {}
    for row in range(5, sheet.max_row + 1):
        code = sheet.cell(row=row, column=2).value
        if code:
            values[str(code)] = Decimal(str(sheet.cell(row=row, column=4).value))
            number_formats[str(code)] = sheet.cell(row=row, column=4).number_format

    assert values["4135"] == D("1000.00")
    assert values["510506"] == D("250.00")
    assert number_formats["4135"] == "#,##0.00"

    totals = {
        str(sheet.cell(row=row, column=1).value): Decimal(str(sheet.cell(row=row, column=4).value))
        for row in range(5, sheet.max_row + 1)
        if str(sheet.cell(row=row, column=1).value or "").startswith("UTILIDAD")
    }
    assert totals["UTILIDAD NETA"] == D("750.00")


def test_export_balance_sheet_xlsx(
    client: TestClient,
    db_session: Session,
) -> None:
    _post_entry(
        db_session,
        date(2026, 3, 2),
        "Ingreso balance",
        [("110505", D("500.00"), D("0.00")), ("4135", D("0.00"), D("500.00"))],
    )

    response = client.get(
        "/api/v1/reports/export/xlsx",
        params={
            "report": "balance-sheet",
            "as_of_date": "2026-03-31",
        },
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert 'filename="balance-general.xlsx"' in response.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    sheet = workbook.active
    assert sheet["A1"].value == "VANTAGE - Balance General"
    assert "Empresa:" in str(sheet["A2"].value)

    rows = {
        str(sheet.cell(row=row, column=1).value): Decimal(str(sheet.cell(row=row, column=4).value))
        for row in range(5, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value
    }
    assert rows["TOTAL ACTIVO"] == D("500.00")
    assert rows["TOTAL PATRIMONIO"] == D("500.00")
    assert rows["PASIVO + PATRIMONIO"] == D("500.00")
